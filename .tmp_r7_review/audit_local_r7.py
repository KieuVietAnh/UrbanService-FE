from __future__ import annotations

import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph


DOCX = Path(r"C:\Users\richdesu\Downloads\Report7_FinalProjectReport.docx")
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "cp": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
}


def iter_block_items(parent):
    parent_elm = parent.element.body if isinstance(parent, DocumentType) else parent._tc
    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


doc = Document(DOCX)
blocks: list[dict] = []
paragraph_index = 0
table_index = 0
for block in iter_block_items(doc):
    if isinstance(block, Paragraph):
        paragraph_index += 1
        text = norm(block.text)
        if text:
            blocks.append(
                {
                    "kind": "p",
                    "location": f"p#{paragraph_index}",
                    "style": block.style.name if block.style else "",
                    "text": text,
                }
            )
    else:
        table_index += 1
        for r_idx, row in enumerate(block.rows, start=1):
            for c_idx, cell in enumerate(row.cells, start=1):
                text = norm(" ".join(p.text for p in cell.paragraphs))
                if text:
                    blocks.append(
                        {
                            "kind": "cell",
                            "location": f"table#{table_index} r{r_idx}c{c_idx}",
                            "style": "",
                            "text": text,
                        }
                    )


def contexts(pattern: str, *, regex: bool = False, flags: int = re.I, limit: int = 20):
    rx = re.compile(pattern if regex else re.escape(pattern), flags)
    out = []
    for item in blocks:
        if rx.search(item["text"]):
            out.append({k: item[k] for k in ("location", "style", "text")})
            if len(out) >= limit:
                break
    return out


print("DOCUMENT SUMMARY")
print(f"path={DOCX}")
print(f"bytes={DOCX.stat().st_size}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print(f"nonempty_blocks={len(blocks)}")

headings = [b for b in blocks if b["kind"] == "p" and b["style"].startswith("Heading")]
print(f"headings={len(headings)} heading_styles={dict(Counter(h['style'] for h in headings))}")

print("\nHEADINGS")
for h in headings:
    print(f"{h['location']}\t{h['style']}\t{h['text']}")

queries = {
    "wrong_ack_names": r"Nguyen Thanh Nam|Vong Tai Dong|Truong Thien Loc|Vo Hoang Tuan Kiet",
    "project_codes": r"SU26SE\d+",
    "phone_verification": r"phone[- ]verification|phone verified|verify (?:the )?phone|SMS OTP|phone OTP",
    "email_verification": r"email[- ]verification|email verified|verify (?:the )?email|email OTP",
    "generic_placeholders": r"\[[^\]]*(?:insert|define|include|descript|project team|to be)[^\]]*\]",
    "project_title": r"Project title\s*:",
    "zalo": r"\bZalo\b",
    "twilio": r"\bTwilio\b",
    "test_count_claims": r"18\s+Normal|14\s+Abnormal|9\s+Boundary|31\s+(?:unit\s+)?tests?|1\s+failed|1\s+blocked",
    "known_glued_words": r"arearesponsible|clientfacing|longrunning|expirationcontrolled|coloronly",
    "cost_na": r"Estimated total cost\s*:\s*N/?A",
    "acknowledgement": r"Acknowledg",
    "workflow_11_12": r"Workflow\s+(?:11|12)\b",
    "provider_roles": r"Service Provider|Provider Coordinator|Service Operator Staff|SERVICEOPERATORSTAFF",
}

print("\nTARGETED CONTEXTS")
for label, pattern in queries.items():
    rows = contexts(pattern, regex=True, limit=40)
    print(f"\n[{label}] count_shown={len(rows)}")
    for row in rows:
        print(f"- {row['location']} | {row['style']} | {row['text']}")

print("\nBRACKETED TEXT")
brackets = []
for item in blocks:
    for match in re.finditer(r"\[[^\]\r\n]{3,250}\]", item["text"]):
        brackets.append((item["location"], match.group(0)))
for location, text in brackets:
    print(f"- {location}: {text}")
print(f"bracketed_count={len(brackets)}")

print("\nNUMBERED HEADING DUPLICATES")
num_map: defaultdict[str, list[dict]] = defaultdict(list)
for h in headings:
    m = re.match(r"^(\d+(?:\.\d+){1,4})\b", h["text"])
    if m:
        num_map[m.group(1)].append(h)
for number, items in sorted(num_map.items()):
    if len(items) > 1:
        print(f"{number}: " + " || ".join(f"{i['location']} {i['text']}" for i in items))

with zipfile.ZipFile(DOCX) as zf:
    names = zf.namelist()
    media = [i for i in zf.infolist() if i.filename.startswith("word/media/")]
    print("\nPACKAGE SUMMARY")
    print(f"zip_parts={len(names)} media_files={len(media)} media_uncompressed_bytes={sum(i.file_size for i in media)}")
    for info in sorted(media, key=lambda i: i.file_size, reverse=True)[:15]:
        print(f"- {info.filename}: {info.file_size}")

    document_xml = zf.read("word/document.xml")
    print("\nREVISION/CONTROL COUNTS")
    for tag in ("ins", "del", "moveFrom", "moveTo", "commentRangeStart", "sdt"):
        print(f"w:{tag}={document_xml.count((b'<w:' + tag.encode('ascii')))}")

    if "docProps/app.xml" in names:
        app = ET.fromstring(zf.read("docProps/app.xml"))
        print("\nAPP PROPERTIES")
        for key in ("Pages", "Words", "Characters", "Paragraphs", "Lines", "Company", "Application"):
            node = app.find(f"cp:{key}", NS)
            print(f"{key}={node.text if node is not None else ''}")

print("\nJSON SUMMARIES")
for json_path in (Path(r"D:\US\.tmp_r7_review\style_lint.json"), Path(r"D:\US\.tmp_r7_review\a11y.json")):
    if not json_path.exists():
        continue
    data = json.loads(json_path.read_text(encoding="utf-8"))
    print(f"{json_path.name}: type={type(data).__name__}")
    if isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, list):
                print(f"- {key}: list[{len(value)}]")
                for sample in value[:8]:
                    print(f"  * {sample}")
            elif isinstance(value, dict):
                print(f"- {key}: dict keys={list(value)[:20]}")
            else:
                print(f"- {key}: {value}")
    else:
        print(f"- length={len(data)}")
