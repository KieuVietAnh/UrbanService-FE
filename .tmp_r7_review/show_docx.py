from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document


parser = argparse.ArgumentParser()
parser.add_argument("--start", type=int)
parser.add_argument("--end", type=int)
parser.add_argument("--query")
parser.add_argument("--tables", action="store_true")
parser.add_argument("--table-index", type=int)
args = parser.parse_args()

doc = Document(Path(r"C:\Users\richdesu\Downloads\Report7_FinalProjectReport.docx"))


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


rx = re.compile(args.query, re.I) if args.query else None
if args.table_index is None:
    for idx, p in enumerate(doc.paragraphs, 1):
        text = clean(p.text)
        if not text:
            continue
        if args.start is not None and idx < args.start:
            continue
        if args.end is not None and idx > args.end:
            continue
        if rx and not rx.search(text):
            continue
        print(f"p#{idx} [{p.style.name if p.style else ''}] {text}")

if args.tables:
    for ti, table in enumerate(doc.tables, 1):
        if args.table_index is not None and ti != args.table_index:
            continue
        for ri, row in enumerate(table.rows, 1):
            cells = [clean(c.text) for c in row.cells]
            joined = " | ".join(cells)
            if rx and not rx.search(joined):
                continue
            print(f"table#{ti} row#{ri}: {joined}")
