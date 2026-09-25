from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document


DOCX = Path(r"C:\Users\richdesu\Downloads\Report7_FinalProjectReport.docx")
doc = Document(DOCX)


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def show_paragraph_range(start: int, end: int, title: str):
    print(f"\n=== {title} (p#{start}-p#{end}) ===")
    for idx in range(start, min(end, len(doc.paragraphs)) + 1):
        p = doc.paragraphs[idx - 1]
        t = clean(p.text)
        if t:
            print(f"p#{idx} [{p.style.name if p.style else ''}] {t}")


def show_matching_rows(label: str, pattern: str):
    rx = re.compile(pattern, re.I)
    print(f"\n=== MATCHING TABLE ROWS: {label} ===")
    count = 0
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            cells = [clean(c.text) for c in row.cells]
            joined = " | ".join(cells)
            if rx.search(joined):
                count += 1
                print(f"table#{ti} row#{ri}: {joined}")
    print(f"count={count}")


print("=== FRONT MATTER AND ACKNOWLEDGEMENT ===")
show_paragraph_range(1, 176, "front matter")

print("\n=== TOC ENTRIES ===")
for idx, p in enumerate(doc.paragraphs, 1):
    style = p.style.name if p.style else ""
    text = clean(p.text)
    if style.lower().startswith("toc") and text:
        print(f"p#{idx} [{style}] {text}")

show_paragraph_range(287, 365, "scope, roles, limitations")
show_paragraph_range(366, 460, "project management")
show_paragraph_range(461, 539, "SRS overview and authentication")
show_paragraph_range(1068, 1151, "NFR and requirement appendix")
show_paragraph_range(1167, 1262, "database and detailed design")
show_paragraph_range(1263, 1363, "testing")
show_paragraph_range(1364, 1470, "release package and account guide")

show_matching_rows("project/team/change history", r"project title|project code|group name|member|version|change|date")
show_matching_rows("actors and role boundaries", r"service operator staff|provider coordinator|service provider|system staff|interaction manager")
show_matching_rows("test metrics", r"31/31|32/32|18\s+normal|14\s+abnormal|9\s+boundary|failed|blocked|passed|test case")
show_matching_rows("database table totals", r"42|50|class table|database table|entity")

with zipfile.ZipFile(DOCX) as zf:
    xml = zf.read("word/document.xml")
    print("\n=== EXACT XML CONTROL COUNTS ===")
    for tag in ("ins", "del", "moveFrom", "moveTo", "instrText", "fldChar", "commentRangeStart"):
        count = len(re.findall(rb"<w:" + tag.encode("ascii") + rb"(?:\s|>)", xml))
        print(f"w:{tag}={count}")
    print(f"comments_part={'word/comments.xml' in zf.namelist()}")
